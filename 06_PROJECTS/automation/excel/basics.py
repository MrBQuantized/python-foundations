import openpyxl
from pathlib import Path

PATH = Path(__file__).resolve().parents[0]
print(PATH)
FILE = PATH / 'myfile.xlsx'


# ============================================
# Reading and Writing to an Existing workbook
# ============================================

# Load the workbook
workbook = openpyxl.load_workbook(FILE)

# Select the active sheet
sheet = workbook.active

# # Set the Sheet's title
# sheet.title = "Changed"

# # values in the sheets cell can be accessed as with a dictionary
# print(sheet['A1'].value)           # Prints the value of column A1
# print(sheet['B1'].value)           # Prints the value of column B1 
# print(sheet['C1'].value)           # Prints the value of column C1
# print(sheet['A2'].value)           # Prints the value of column A2
# print(sheet['B2'].value)           # Prints the value of column B2
# print(sheet['C2'].value)           # Prints the value of column C2
# print(sheet['A3'].value)           # Prints the value of column A3
# print(sheet['B3'].value)           # Prints the value of column B3
# print(sheet['C3'].value)           # Prints the value of column C3

# # We can set values to cells as well
# sheet['A3'] = 70 
# sheet['B3'] = 80 
# sheet['C3'] = 90 
# sheet['A5'] = True 
# sheet['B5'] = False 
# sheet['C5'] = 'Exception'
# sheet['F6'] = "=B2+C2"              # We can even pass in formulas as strings.
# # save the changes made to the workbook
# workbook.save(FILE)

# ============================================
# Creating New Sheets
# ============================================
# workbook.create_sheet("New_sheet")
# sheet = workbook["New_sheet"]

# sheet['A1'] = "Name"
# sheet['B1'] = "Age"
# sheet['C1'] = "Class"

# workbook.save(FILE)

# ============================================
# Creating New Workbook
# ============================================

NEW_WORKBOOK = PATH / 'new_workbook.xlsx'

workbook_2 = openpyxl.Workbook()
sheet = workbook_2.active

sheet['A1'] = 10   
sheet['B1'] = 20   
sheet['C1'] = 30   
sheet['A2'] = 40   
sheet['B2'] = 50   
sheet['C2'] = 60   
sheet['A3'] = 70   
sheet['B3'] = 80   
sheet['C3'] = 90  

# # Inserting rows and columns
# sheet.insert_rows(1)            # inserts a row at index position 1
# sheet.insert_cols(1)            # inserts a column at index position 1

# # We can also change the amount
# sheet.insert_rows(1, amount=5)           
# sheet.insert_cols(1, amount=3)           


# Merging cells
# sheet.merge_cells("A1:C3")
# sheet['A1']  = "Hello World!"

# Moving a range of cells
# sheet.move_range("A1:C3", rows=6, cols=4 )    # Cells A1 to C3 get moved, 6 rows down and 4 columns to the right
# # we can shift the range of cells to the left as well
sheet.move_range("C1:C3", rows=0, cols=-1 )    
workbook_2.save(NEW_WORKBOOK)